"""Draft network-wide SKU allocation matrix (read only). No stock claims or database writes."""
from __future__ import annotations
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
import math
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

MODEL = Path(__file__).resolve().parent / 'catboost_first_model.cbm'
SHELF = {'второе':5,'вторые блюда':5,'салат':3,'салаты':3,'завтрак':3,'завтраки':3,
         'завтарак':3,'сэндвич':3,'сэндвичи':3,'суп':3,'супы':3,'япония':2,
         'напиток':7,'напитки':7}
CATEGORIES = {'второе':'Вторые блюда','вторые блюда':'Вторые блюда','салат':'Салаты',
 'салаты':'Салаты','завтрак':'Завтраки','завтраки':'Завтраки','завтарак':'Завтраки',
 'сэндвич':'Сэндвичи','сэндвичи':'Сэндвичи','суп':'Супы','супы':'Супы',
 'япония':'Япония','напиток':'Напитки','напитки':'Напитки'}

@st.cache_resource(show_spinner=False)
def model_load(path, modified):
    from catboost import CatBoostRegressor
    model = CatBoostRegressor()
    model.load_model(path)
    expected = ['Точка','SKU','Категория','weekday','week_parity','month','day_of_month','Срок, дней']
    if list(model.feature_names_) != expected:
        raise RuntimeError('Версия модели не соответствует набору признаков')
    return model

@st.cache_data(ttl=300, show_spinner=False)
def plan_for_date(_pg_connection, menu_date):
    with _pg_connection() as conn:
        with conn.cursor() as cur:
            cur.execute('''
              SELECT point_number, NULLIF(LTRIM(TRIM(sku::text),'0'),'') AS sku,
                     MAX(product_name) AS product_name, MAX(category) AS category,
                     SUM(fact_plan_qty)::numeric AS plan_qty
              FROM planning.menu_plan
              WHERE menu_date=%s AND point_number BETWEEN 1 AND 40
                AND sku IS NOT NULL AND fact_plan_qty > 0
              GROUP BY point_number, NULLIF(LTRIM(TRIM(sku::text),'0'),'')
              ORDER BY sku, point_number
            ''',(menu_date,))
            rows=cur.fetchall(); cols=[x.name for x in cur.description]
    return pd.DataFrame(rows,columns=cols)

def category_warnings(frame):
    names=frame['product_name'].fillna('').str.casefold()
    cats=frame['category'].fillna('').str.casefold()
    wrong=(names.str.contains('морс|компот',regex=True) & ~cats.isin(['напитки','напиток'])) | \
          (names.str.contains('онигири',regex=False) & cats.ne('япония'))
    return wrong

def compute_matrix(plan, menu_date, pct=30):
    d=plan.copy()
    if d.empty:return d
    d['SKU']=d['sku'].astype(str)
    d['Точка']='Т'+d['point_number'].astype(int).astype(str)
    d['Категория']=d['category'].fillna('').astype(str).str.strip()
    d['Проверить категорию']=category_warnings(d)
    d['Срок, дней']=d['Категория'].str.casefold().map(SHELF)
    d['Текущий план']=pd.to_numeric(d['plan_qty'],errors='coerce').fillna(0).astype(int)
    d['Прогноз продаж']=pd.NA
    d['Рекомендуется']=d['Текущий план'].copy()
    d['Причина']='Неизвестный срок категории: оставлен план'
    usable=(d['Срок, дней'].notna() & ~d['Проверить категорию'])
    if usable.any():
        if not MODEL.is_file():
            raise FileNotFoundError('Модель catboost_first_model.cbm не найдена рядом с ml_network.py')
        m=model_load(str(MODEL),MODEL.stat().st_mtime_ns)
        x=d.loc[usable,['Точка','SKU','Категория','Срок, дней']].copy()
        x['weekday']=int(menu_date.weekday())
        x['week_parity']=str(menu_date.isocalendar().week%2)
        x['month']=int(menu_date.month)
        x['day_of_month']=int(menu_date.day)
        x['Срок, дней']=x['Срок, дней'].astype(int)
        order=['Точка','SKU','Категория','weekday','week_parity','month','day_of_month','Срок, дней']
        pred=pd.Series(m.predict(x[order]),index=x.index).clip(lower=0)
        d.loc[usable,'Прогноз продаж']=pred.round(2)
        base=d.loc[usable,'Текущий план']
        lo=(base*(1-pct/100)).apply(math.floor).clip(lower=0)
        hi=(base*(1+pct/100)).apply(math.ceil)
        rec=pred.apply(math.ceil).clip(lower=lo,upper=hi).astype(int)
        d.loc[usable,'Рекомендуется']=rec
        d.loc[usable,'Причина']='CatBoost, изменение ограничено +/- '+str(pct)+'%; без подтверждённого остатка'
    d.loc[d['Проверить категорию'],'Причина']='Конфликт наименования/категории: оставлен план'
    d['Рекомендуется']=d['Рекомендуется'].astype(int)
    d['Изменение']=d['Рекомендуется']-d['Текущий план']
    return d

def exclusion_rules(_pg_connection):
    with _pg_connection() as conn:
        with conn.cursor() as c:
            c.execute("SELECT point_code,scope,pattern FROM ml.supply_rules WHERE is_active=TRUE AND action='exclude'")
            return c.fetchall()

def apply_exclusions(d, rules):
    d=d.copy()
    for point,scope,pattern in rules:
        term=str(pattern or '').strip().casefold()
        if not term:continue
        mask=d['Точка'].eq(str(point).replace('T','Т'))
        if scope=='sku':mask &= d['SKU'].str.lstrip('0').eq(term.lstrip('0'))
        elif scope=='category': mask &= d['Категория'].str.casefold().eq(term)
        elif scope=='name_contains': mask &= d['product_name'].fillna('').str.casefold().str.contains(term,regex=False)
        else:continue
        d.loc[mask,'Рекомендуется']=0
        d.loc[mask,'Причина']='Подтверждённый запрет'
    d['Изменение']=d['Рекомендуется']-d['Текущий план']
    return d

def excel_matrix(detail, menu_date):
    points=[f'Т{i}' for i in range(1,41)]
    grid=detail.pivot_table(index=['SKU','product_name','Категория'],columns='Точка',values='Утверждаемое количество',aggfunc='first',fill_value=0)
    grid=grid.reindex(columns=points,fill_value=0).reset_index()
    grid['ИТОГО']=grid[points].sum(axis=1)
    w=Workbook(); ws=w.active;ws.title='Матрица ИИ'
    ws.append(['SKU','Название блюда','Категория',*points,'ИТОГО'])
    for row in grid.itertuples(index=False,name=None):ws.append(list(row))
    ws.append(['ИТОГО','','',*[int(grid[x].sum()) for x in points],int(grid['ИТОГО'].sum())])
    ws.freeze_panes='D2';ws.auto_filter.ref=f'A1:{get_column_letter(4+len(points))}{ws.max_row-1}'
    ws.column_dimensions['A'].width=12;ws.column_dimensions['B'].width=43;ws.column_dimensions['C'].width=20
    for j in range(4,45):ws.column_dimensions[get_column_letter(j)].width=9
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='20344B')
    for cell in ws[ws.max_row]:cell.font=Font(bold=True)
    ws2=w.create_sheet('Проверка изменений')
    cols=['Точка','SKU','product_name','Категория','Текущий план','Прогноз продаж','Рекомендуется','Утверждаемое количество','Изменение','Причина']
    ws2.append(cols)
    for r in detail[cols].itertuples(index=False,name=None):ws2.append([None if pd.isna(v) else v for v in r])
    ws2.freeze_panes='A2';ws2.auto_filter.ref=ws2.dimensions
    for cell in ws2[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='20344B')
    ws2.column_dimensions['C'].width=46;ws2.column_dimensions['J'].width=65
    output=BytesIO();w.save(output);return output.getvalue()

def render_network(pg_connection, today:date):
    st.subheader('ИИ — матрица рекомендуемых поставок')
    st.caption('Одна матрица по всем точкам. Источники те же, что у DataLens: PostgreSQL. Реализация и остатки остаются в DataLens.')
    day=st.date_input('Дата меню',value=today,key='ml_network_day')
    with st.expander('Настройки первого пилота',expanded=False):
        cap=st.slider('Ограничение изменения от текущего плана, %',10,50,30,5,key='ml_network_cap')
        st.warning('Это черновой прогноз наблюдаемых продаж, НЕ оптимальная поставка. Физические остатки и время распродажи пока не учтены. Цветовые минимумы не применены без цветов точек.')
    if st.button('Сформировать матрицу всей сети',type='primary',key='ml_network_build'):
        try:
            with st.spinner('Читаю выбранное меню и рассчитываю матрицу...'):
                plan=plan_for_date(pg_connection,day)
                if plan.empty:
                    st.warning('Для даты нет строк в planning.menu_plan. Выбери дату, уже загруженную в PostgreSQL.')
                    st.session_state.pop('ml_network_result',None)
                    return
                result=compute_matrix(plan,day,cap)
                try: result=apply_exclusions(result,exclusion_rules(pg_connection))
                except Exception as exc: st.warning(f'Запреты не удалось прочитать, выдача отключена: {exc}');return
                result['Утверждаемое количество']=result['Рекомендуется']
                st.session_state['ml_network_result']=(day,cap,result)
        except Exception as exc:
            st.error(f'Не удалось рассчитать матрицу: {exc}')
    state=st.session_state.get('ml_network_result')
    if not state or state[0]!=day or state[1]!=cap:return
    result=state[2].copy()
    c1,c2,c3=st.columns(3)
    c1.metric('Товарных назначений',len(result))
    c2.metric('Сейчас по плану',int(result['Текущий план'].sum()))
    c3.metric('Черновая рекомендация',int(result['Рекомендуется'].sum()))
    issues=int(result['Проверить категорию'].sum())
    if issues:st.warning(f'{issues} позиций с подозрительной категорией — по ним оставлен текущий план.')
    st.markdown('#### Матрица SKU × Т1–Т40')
    pts=[f'Т{i}' for i in range(1,41)]
    pivot=result.pivot_table(index=['SKU','product_name'],columns='Точка',values='Рекомендуется',aggfunc='first',fill_value=0).reindex(columns=pts,fill_value=0).reset_index()
    pivot['ИТОГО']=pivot[pts].sum(axis=1)
    st.dataframe(pivot,hide_index=True,use_container_width=True,height=420)
    st.markdown('#### Проверка и ручная корректировка')
    st.caption('Редактируй «Утверждаемое количество». Изменения действуют только в скачанном файле, PostgreSQL не изменяется.')
    show=['Точка','SKU','product_name','Категория','Текущий план','Прогноз продаж','Рекомендуется','Утверждаемое количество','Причина']
    edited=st.data_editor(result[show],disabled=[c for c in show if c!='Утверждаемое количество'],hide_index=True,use_container_width=True,
        key=f'ml_network_edit_{day}_{cap}',column_config={'Утверждаемое количество':st.column_config.NumberColumn('Утверждаемое количество',min_value=0,step=1,required=True)})
    final=result.copy();final['Утверждаемое количество']=pd.to_numeric(edited['Утверждаемое количество'],errors='coerce').fillna(final['Рекомендуется']).clip(lower=0).round().astype(int)
    st.download_button('Скачать матрицу Excel',data=excel_matrix(final,day),file_name=f'AI_supply_matrix_{day}.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',type='primary')
    st.caption('Файл — черновик для менеджера, не автоматическое изменение плана. Проверяй ограничения по цветам, категориям и реальным остаткам до утверждения.')
